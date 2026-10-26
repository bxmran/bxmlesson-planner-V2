import unittest, sys, tempfile, json
from pathlib import Path
from docx import Document
sys.path.insert(0,str(Path(__file__).resolve().parents[1]/'scripts'))
from parse_schedule import extract
from render_school_docx import generate, validate_draft, INDEX,ACTIVITY_FIELDS,ASSESS_FIELDS

class SchoolSkillTests(unittest.TestCase):
    def sample_lesson(self,n):
        obj={'number':n,'hours':'2','form':'','location':'','design_thought':{'intro':'采用分层任务驱动教学。','before':'预习核心概念','during':'讲解并完成任务','after':'作业反馈修订'}}
        for name in INDEX:
            if name not in obj:
                obj[name]=f'第{n}次课 {name}：学生完成对应问题分析，并由教师依据材料给予针对性反馈。'
        # Quality tests use approximately 400 characters across teaching + learning,
        # and around 150 characters for pre-class tasks.
        obj['teaching_methods']='\n'.join([
            f'1. 情境案例法：围绕第{n}次课重点出示两种真实业务情境，引导学生比较差异、标注关键信息，再用典型问题追问判断依据；教师纠正把现象等同结论的误区，并提示使用同一标准核验。',
            '2. 图示框架法：按课堂知识链绘制条件对比图，结合教材中的核心概念分步讲解，让学生填入案例要素、补齐推理过程；对混淆的概念通过反例核查并反馈。',
            '3. 任务驱动法：布置一页分析工作单，先独立作答再小组比较，对每项判断写出材料来源；教师巡视诊断常见错误，选取典型回答展示并组织订正。'
        ])
        obj['learning_methods']='\n'.join([
            '1. 自主学习：按教材节次整理关键词，先独立填写判断清单，对不确定的信息做出标记，并提出待解释的问题。',
            '2. 小组探究：分工收集课堂案例证据，合作完成要素比较表，各自说明判断理由，依据教师提供的标准互相核对。',
            '3. 迁移订正：将方法用于另一简短情境，记录讨论前后的结论变化与修正依据，形成个人学习记录。'
        ])
        obj['preparation']='\n'.join([
            f'1. 阅读第{n}次课对应教材节次，圈画概念关键词，尝试用一句话解释其中三个重要术语。',
            '2. 观察本次主题相关企业或业务情境，记录两项可核实的事实及来源，不能确认的事项标记待查。',
            '3. 回答“关键判断应依靠哪些条件”并写出初步理由，把概念清单、案例记录和疑问整理成一页预习卡，上课时带来供分组核对。'
        ])
        phases=['情境导入','新知讲解','案例讨论','任务操作','总结反馈']
        def row(phase,time):return {'phase':phase,'time':time,'content':'围绕本次课核心内容开展任务','method':'任务驱动法','teacher':'提出问题并指导纠错','student':'分组整理材料并展示','ethics':'诚信与证据意识','goals':'目标1、目标2'}
        obj['implementation']={'before':row('课前预习','课前'),
            'during':[row(p,f'{tm}分钟') for p,tm in zip(phases,[10,30,25,20,5])], 'after':row('课后反馈','课后')}
        obj['assessment']=[{'module':name,'weight':weight,'evidence':'工作单','criteria':'判断有依据','evaluator':'教师','feedback':'批注'} for name,weight in [('课前','20%'),('课中','50%'),('课后','30%')]]
        return obj

    def test_school_template_is_lesson_only(self):
        src=Path(__file__).resolve().parents[1]/'assets/school_lesson_only_blank.docx'
        d=Document(src)
        self.assertEqual(len(d.tables),1)
        self.assertEqual(len(d.tables[0].rows),49)
        self.assertEqual(d.tables[0].rows[1].cells[0].text,'授课章节')
        self.assertFalse(any('首页' in p.text for p in d.paragraphs))

    def test_dynamic_22_lessons_and_blank_unknown_info(self):
        with tempfile.TemporaryDirectory() as td:
            draft=Path(td)/'draft.json';output=Path(td)/'result.docx'
            # Every number must be supported, not hard-coded to 16 or 20.
            draft.write_text(json.dumps({'lessons':[self.sample_lesson(n) for n in range(1,23)]},ensure_ascii=False),encoding='utf8')
            result=generate(draft,output)
            self.assertEqual(result['count'],22)
            out=Document(output)
            self.assertEqual(len(out.tables),22)
            self.assertEqual(out.tables[-1].rows[2].cells[8].text,'')
            self.assertEqual(out.tables[0].rows[1].cells[8].text,'')
            self.assertIn('第22次课',[p.text for p in out.paragraphs])

    def test_parse_school_schedule(self):
        with tempfile.TemporaryDirectory() as td:
            p=Path(td)/'schedule.docx'
            d=Document();t=d.add_table(rows=1,cols=8)
            for idx,val in {0:'课次',1:'学时数',2:'授课内容',6:'作业布置',7:'备注'}.items():t.rows[0].cells[idx].text=val
            for n in range(1,26):
                row=t.add_row();row.cells[0].text=str(n);row.cells[1].text='2';row.cells[2].text=f'第{n}讲';row.cells[6].text=f'作业{n}'
            d.save(p)
            self.assertEqual(len(extract(p)['lessons']),25)
            self.assertEqual(len(extract(p,'1-16')['lessons']),16)
            self.assertEqual(len(extract(p,'21-25')['lessons']),5)

    def test_shallow_method_and_preparation_rejected(self):
        for field, shallow in [('teaching_methods','讲授法、案例法、讨论法。'),
                                ('learning_methods','自主学习、合作学习、课后反思。'),
                                ('preparation','阅读教材、思考问题。')]:
            lesson=self.sample_lesson(1)
            lesson[field]=shallow
            with self.subTest(field=field),self.assertRaisesRegex(ValueError,'需至少'):
                validate_draft({'lessons':[lesson]})

    def test_numbered_items_enforced(self):
        lesson=self.sample_lesson(1)
        lesson['teaching_methods']='教学法可操作性与成果反馈。'*18
        with self.assertRaisesRegex(ValueError,'编号的针对性任务'):
            validate_draft({'lessons':[lesson]})

    def test_quality_fields_good(self):
        lesson=self.sample_lesson(1)
        self.assertEqual(validate_draft({'lessons':[lesson]})['count'],1)

    def test_length_limits_and_readable_spacing(self):
        lesson=self.sample_lesson(1)
        with tempfile.TemporaryDirectory() as td:
            p=Path(td)/'d.json';out=Path(td)/'o.docx'
            p.write_text(json.dumps({'lessons':[lesson]},ensure_ascii=False),encoding='utf8')
            generate(p,out)
            d=Document(out)
            for i,j in [(14,3),(21,3),(22,3),(26,2),(33,3),(33,5),(44,0),(46,3)]:
                par=d.tables[0].rows[i].cells[j].paragraphs[0]
                self.assertEqual(par.paragraph_format.line_spacing.pt,20)
            methods=d.tables[0].rows[21].cells[3]
            self.assertEqual(methods.paragraphs[0].paragraph_format.space_after.pt,2)
            self.assertEqual(len(d.tables),1)
        for field, overlong in [('teaching_methods',lesson['teaching_methods']+'教学细节'*80),
                                ('learning_methods',lesson['learning_methods']+'案例复核'*40),
                                ('preparation',lesson['preparation']+'阅读更多材料'*30)]:
            l=self.sample_lesson(1);l[field]=overlong
            with self.subTest(field=field),self.assertRaisesRegex(ValueError,'超过精练上限'):
                validate_draft({'lessons':[l]})

    def test_missing_field_rejected(self):
        lesson=self.sample_lesson(1)
        lesson['reflection_achievement']=''
        with self.assertRaises(ValueError):validate_draft({'lessons':[lesson]})

if __name__=='__main__': unittest.main()
